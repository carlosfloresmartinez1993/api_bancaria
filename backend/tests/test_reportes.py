import io
from unittest.mock import patch

from openpyxl import load_workbook

from app.services.cierre import ejecutar_cierre
from tests.helpers import capturar, crear_cliente, crear_empresa, crear_proyecto, dia, metodo, registrar_salida


def _datos(client, h):
    """Apple → clientes 1234 (P1 2 %, P2 3 %) y 5678 (P1 9 %), con entradas y salidas."""
    empresa = crear_empresa(client, h, "Apple")
    c1 = crear_cliente(client, h, empresa["id"], "1234")
    c2 = crear_cliente(client, h, empresa["id"], "5678")
    p1 = crear_proyecto(client, h, c1["id"], "P1", porcentaje="2")
    p2 = crear_proyecto(client, h, c1["id"], "P2", porcentaje="3")
    p3 = crear_proyecto(client, h, c2["id"], "P1", porcentaje="9")
    capturar(client, h, p1["id"], "1000", fecha=dia(-3))                       # neto 980
    capturar(client, h, p1["id"], "500", fecha=dia(-1), requiere_factura=True)  # neto 490
    capturar(client, h, p2["id"], "100", fecha=dia(-1), metodo_pago_id=metodo("Efectivo"))  # neto 97
    capturar(client, h, p3["id"], "200", fecha=dia(-1))                        # neto 182
    registrar_salida(client, h, p1["id"], "400", "Renta local", fecha=dia(-2))
    registrar_salida(client, h, p3["id"], "50", "Gasolina", fecha=dia(-1), metodo_pago_id=metodo("Efectivo"))
    return empresa, (c1, c2), (p1, p2, p3)


def _por(filas, **claves):
    return next(f for f in filas if all(f[k] == v for k, v in claves.items()))


def test_constructor_por_proyecto_cliente_empresa_y_metodo(client, juan):
    h = juan[1]
    empresa, _, _ = _datos(client, h)
    base = {"empresa_id": empresa["id"], "desde": dia(-5), "hasta": dia(0)}

    r = client.get("/reportes/constructor", params={**base, "agrupar": "proyecto"}, headers=h).json()
    assert r["elaborado_por"] == "Juan Prueba"
    assert [c["clave"] for c in r["columnas"]][:3] == ["empresa", "cliente", "proyecto"]
    p1 = _por(r["filas"], cliente="1234", proyecto="P1")
    assert p1 == {**p1, "movimientos": 2, "monto_bruto": "1500.00", "comision": "30.00", "monto_neto": "1470.00",
                  "num_salidas": 1, "salidas": "400.00", "saldo": "1070.00"}
    assert _por(r["filas"], cliente="5678", proyecto="P1")["saldo"] == "132.00"
    assert r["totales"]["monto_neto"] == "1749.00" and r["totales"]["saldo"] == "1299.00"

    r = client.get("/reportes/constructor", params={**base, "agrupar": "cliente"}, headers=h).json()
    assert {f["cliente"]: f["monto_neto"] for f in r["filas"]} == {"1234": "1567.00", "5678": "182.00"}

    r = client.get("/reportes/constructor", params={**base, "agrupar": "empresa"}, headers=h).json()
    assert len(r["filas"]) == 1 and r["filas"][0]["comision"] == "51.00"

    r = client.get("/reportes/constructor", params={**base, "agrupar": "metodo_pago"}, headers=h).json()
    efectivo = _por(r["filas"], metodo_pago="Efectivo")
    assert efectivo["monto_neto"] == "97.00" and efectivo["salidas"] == "50.00"


def test_constructor_con_seleccion_de_clientes_proyectos_y_contenido(client, juan):
    h = juan[1]
    empresa, (c1, _), (p1, p2, _) = _datos(client, h)
    base = {"empresa_id": empresa["id"], "desde": dia(-5), "hasta": dia(0)}

    r = client.get("/reportes/constructor", params={**base, "terminal_id": [c1["id"]]}, headers=h).json()
    assert {f["cliente"] for f in r["filas"]} == {"1234"}

    r = client.get("/reportes/constructor", params={**base, "proyecto_id": [p1["id"], p2["id"]],
                                                    "contenido": "entradas"}, headers=h).json()
    assert len(r["filas"]) == 2 and "saldo" not in r["totales"] and "salidas" not in r["totales"]

    r = client.get("/reportes/constructor", params={**base, "contenido": "salidas", "texto": "renta"}, headers=h).json()
    assert len(r["filas"]) == 1 and r["totales"]["salidas"] == "400.00"

    r = client.get("/reportes/constructor", params={**base, "contenido": "entradas", "requiere_factura": True,
                                                    "agrupar": "detalle"}, headers=h).json()
    assert len(r["filas"]) == 1 and r["filas"][0]["factura"] == "Sí" and r["filas"][0]["monto_bruto"] == "500.00"

    r = client.get("/reportes/constructor", params={**base, "agrupar": "detalle"}, headers=h).json()
    assert [f["tipo"] for f in r["filas"]].count("Salida") == 2 and len(r["filas"]) == 6
    assert r["filas"][0]["registro"] == "Juan Prueba"


def test_estado_de_cuenta_con_saldo_corrido(client, juan):
    empresa, (c1, _), _ = _datos(client, juan[1])
    r = client.get("/reportes/estado-cuenta", headers=juan[1],
                   params={"empresa_id": empresa["id"], "desde": dia(-3), "hasta": dia(0)})
    assert r.status_code == 200
    saldos = [f["saldo"] for f in r.json()["filas"]]
    assert saldos[:3] == ["0.00", "980.00", "580.00"] and saldos[-1] == "1299.00"
    assert r.json()["parametros"]["saldo_final"] == "1299.00"

    # Solo un cliente
    r = client.get("/reportes/estado-cuenta", headers=juan[1],
                   params={"empresa_id": empresa["id"], "terminal_id": [c1["id"]], "desde": dia(-3), "hasta": dia(0)})
    assert r.json()["parametros"]["saldo_final"] == "1167.00"


def test_exportaciones_incluyen_elaborado_por(client, admin, juan):
    empresa, _, _ = _datos(client, juan[1])
    params = {"empresa_id": empresa["id"], "desde": dia(-5), "hasta": dia(0), "agrupar": "detalle"}
    r = client.get("/reportes/constructor", params={**params, "formato": "csv"}, headers=juan[1])
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    texto = r.content.decode("utf-8-sig")
    assert "Bruto" in texto and "Elaborado por: Juan Prueba" in texto

    r = client.get("/reportes/constructor", params={**params, "formato": "xlsx"}, headers=juan[1])
    hoja = load_workbook(io.BytesIO(r.content)).active
    assert hoja.cell(row=hoja.max_row, column=1).value.startswith("Elaborado por: Juan Prueba")
    assert hoja.max_row == 1 + 6 + 1 + 2  # encabezado + 6 registros + total + línea en blanco y pie

    r = client.get("/reportes/saldos", params={"formato": "pdf"}, headers=admin[1])
    assert r.content.startswith(b"%PDF")

    for ruta in ("auditoria-captura", "saldos", "constructor"):
        assert client.get(f"/reportes/{ruta}", headers=admin[1]).status_code == 200
    assert client.get("/reportes/resumen-mensual", params={"anio": 2026}, headers=admin[1]).status_code == 200
    r = client.get("/reportes/conciliacion-diaria", params={"fecha": dia(-1)}, headers=juan[1])
    assert r.status_code == 200 and len(r.json()["filas"]) == 3


def test_saldos_incluyen_comisiones(client, juan):
    _datos(client, juan[1])
    fila = client.get("/reportes/saldos", headers=juan[1]).json()["filas"][0]
    assert fila == {**fila, "empresa": "Apple", "ingresos_brutos": "1800.00", "comisiones": "51.00",
                    "ingresos_netos": "1749.00", "salidas": "450.00", "saldo": "1299.00"}


def test_cierre_diario_registra_error_si_no_hay_smtp_y_envio_exitoso(client, admin, juan):
    _datos(client, juan[1])
    assert ejecutar_cierre()["estado"] == "error"

    with patch("app.services.cierre.enviar_correo") as enviar:
        resultado = ejecutar_cierre(forzar=True)
    assert resultado["estado"] == "enviado"
    destinatarios, asunto = enviar.call_args.args[:2]
    assert destinatarios == ["admin@test.com"] and "Cierre diario" in asunto

    with patch("app.services.cierre.enviar_correo") as enviar:
        assert ejecutar_cierre()["estado"] == "ya_enviado"
        enviar.assert_not_called()

    acciones = [f["accion"] for f in client.get("/reportes/bitacora", headers=admin[1]).json()["filas"]]
    assert "CIERRE_ERROR" in acciones and "CIERRE_ENVIADO" in acciones
