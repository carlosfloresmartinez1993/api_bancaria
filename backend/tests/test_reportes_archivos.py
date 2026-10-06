import io
from unittest.mock import patch

from openpyxl import load_workbook

from app.services.cierre import ejecutar_cierre
from tests.helpers import capturar, crear_empresa, crear_terminal, dia


def _datos(client, h):
    empresa = crear_empresa(client, h)
    terminal = crear_terminal(client, h, empresa["id"], porcentaje="2")
    capturar(client, h, terminal["id"], "1000", fecha=dia(-3))
    capturar(client, h, terminal["id"], "500", fecha=dia(-1))
    client.post("/salidas", headers=h, json={"empresa_id": empresa["id"], "monto": "400", "destino": "Renta local",
                                             "fecha": dia(-2)})
    return empresa, terminal


def test_estado_de_cuenta_con_saldo_corrido(client, juan):
    empresa, _ = _datos(client, juan[1])
    r = client.get("/reportes/estado-cuenta", headers=juan[1],
                   params={"empresa_id": empresa["id"], "desde": dia(-3), "hasta": dia(0)})
    assert r.status_code == 200
    saldos = [f["saldo"] for f in r.json()["filas"]]
    assert saldos == ["0.00", "980.00", "580.00", "1070.00"]
    assert r.json()["parametros"]["saldo_final"] == "1070.00"


def test_busqueda_de_salidas_por_texto(client, juan):
    _datos(client, juan[1])
    r = client.get("/reportes/salidas", params={"texto": "renta"}, headers=juan[1])
    assert len(r.json()["filas"]) == 1
    r = client.get("/reportes/salidas", params={"texto": "%"}, headers=juan[1])
    assert r.json()["filas"] == []  # los comodines se escapan


def test_exportaciones(client, admin, juan):
    empresa, _ = _datos(client, juan[1])
    params = {"empresa_id": empresa["id"], "desde": dia(-5), "hasta": dia(0)}
    r = client.get("/reportes/movimientos-por-empresa", params={**params, "formato": "csv"}, headers=juan[1])
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    assert "Bruto" in r.content.decode("utf-8-sig")

    r = client.get("/reportes/movimientos-por-empresa", params={**params, "formato": "xlsx"}, headers=juan[1])
    hoja = load_workbook(io.BytesIO(r.content)).active
    assert hoja.max_row == 4  # encabezado + 2 movimientos + total

    r = client.get("/reportes/comisiones-por-terminal", params={"formato": "pdf"}, headers=admin[1])
    assert r.content.startswith(b"%PDF")

    for ruta in ("totales-por-empresa", "auditoria-captura", "saldos"):
        assert client.get(f"/reportes/{ruta}", headers=admin[1]).status_code == 200
    assert client.get("/reportes/resumen-mensual", params={"anio": 2026}, headers=admin[1]).status_code == 200
    assert client.get("/reportes/conciliacion-diaria", params={"fecha": dia(-1)}, headers=juan[1]).status_code == 200


def test_cierre_diario_registra_error_si_no_hay_smtp_y_envio_exitoso(client, admin, juan):
    _datos(client, juan[1])
    assert ejecutar_cierre()["estado"] == "error"

    with patch("app.services.cierre.enviar_correo") as enviar:
        resultado = ejecutar_cierre(forzar=True)
    assert resultado["estado"] == "enviado"
    destinatarios, asunto = enviar.call_args.args[:2]
    assert destinatarios == ["admin@test.com"] and "Cierre diario" in asunto

    with patch("app.services.cierre.enviar_correo") as enviar:
        assert ejecutar_cierre()["estado"] == "ya_enviado"
        enviar.assert_not_called()

    acciones = [f["accion"] for f in client.get("/reportes/bitacora", headers=admin[1]).json()["filas"]]
    assert "CIERRE_ERROR" in acciones and "CIERRE_ENVIADO" in acciones


def test_archivos_validan_contenido_y_acceso(client, juan, pedro):
    empresa = crear_empresa(client, juan[1])
    url = f"/empresas/{empresa['id']}/archivos/pdf1"

    falso = {"archivo": ("doc.pdf", b"MZ esto es un ejecutable", "application/pdf")}
    assert client.put(url, files=falso, headers=juan[1]).status_code == 422

    pdf = {"archivo": ("csf.pdf", b"%PDF-1.4 contenido", "application/pdf")}
    r = client.put(url, files=pdf, headers=juan[1])
    assert r.status_code == 200 and r.json()["archivos"]["pdf1"] == url

    assert client.get(url, headers=juan[1]).content.startswith(b"%PDF")
    assert client.get(url, headers=pedro[1]).status_code == 404

    png = {"archivo": ("logo.png", b"\x89PNG\r\n\x1a\n" + b"0" * 20, "image/png")}
    assert client.put(f"/empresas/{empresa['id']}/archivos/logo", files=png, headers=juan[1]).status_code == 200
    assert client.put(f"/empresas/{empresa['id']}/archivos/pdf2", files=png, headers=juan[1]).status_code == 422
